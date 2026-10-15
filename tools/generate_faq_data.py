#!/usr/bin/env python3
"""
Regenerates the two FAQ-derived data files the bot ships with:

  src/data/faqReference.json       used by /community faqgap
  src/data/prefilterVocabulary.json  used to decide which Discord messages are
                                    worth sending to the AI

Usage:
  python3 tools/generate_faq_data.py path/to/faqworkbook.xlsx

Run this whenever the FAQ changes, then commit the two JSON files. Requires
only openpyxl (pip install openpyxl). The bot itself never reads the workbook.
"""
import collections
import html
import json
import os
import re
import sys

import openpyxl

OUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "src", "data")

STOPWORDS = set("""
a an the is are was were be been being do does did will would can could should shall may might must have has had
i you he she it we they what which who whom this that these those am of in on at to for with by from up down out
over under again further then once here there when where why how all any both each few more most other some such
no nor not only own same so than too very just and or but if because as until while about against between into
through during before after above below my your his her its our their yours theirs mine me us them get gets got
also still even ever much many need needs want like make makes made use used using one two per via yet whether
also able within without upon onto across along around among off per etc
""".split())

# Informal words players type that never appear in the polished FAQ text,
# mapped from the formal term that does. Applied only when the formal term
# is present in an entry (whole-word match), so terms stay relevant per entry.
PLAYER_SLANG = {
    "ranked": ["rank", "ranked mode", "comp"],
    "ranked points": ["rp", "rank points", "elo", "mmr"],
    "matchmaking": ["mm", "unfair matches", "gear gap"],
    "hero stones": ["hero shards"],
    "emeralds": ["gems", "premium currency"],
    "gear": ["items", "equipment", "build"],
    "spells": ["active ability", "summoner spell"],
    "ancients": ["passives", "relics"],
    "blitz pass": ["battle pass", "season pass"],
    "login": ["log in", "sign in", "cant login"],
    "refund": ["money back", "wrong purchase", "accidental purchase"],
    "bots": ["ai players", "fake players"],
    "cheating": ["cheater", "hacker", "hacking", "exploiter"],
    "friends": ["add friend", "friend list"],
    "daily rewards": ["daily login", "login bonus"],
    "lucky spin": ["wheel", "gacha", "spin"],
    "chests": ["loot boxes", "crates"],
    "crashes": ["crashing", "force close"],
    "battery": ["battery drain", "overheating"],
    "ban": ["banned", "suspended"],
    "controls": ["unresponsive controls"],
    "abilities": ["skills"],
    "demotion": ["demoted", "rank drop"],
}


def strip_html(raw):
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", html.unescape(str(raw or "")))).strip()


def tokens(text):
    return re.findall(r"[^\W\d_]{3,}", text.lower())


def sheet_rows(ws):
    for r in range(2, ws.max_row + 1):
        q = ws.cell(r, 3).value
        if not q:
            continue
        sims = [str(ws.cell(r, c).value) for c in range(10, ws.max_column + 1) if ws.cell(r, c).value]
        yield {
            "id": str(ws.cell(r, 2).value),
            "category": ws.cell(r, 1).value,
            "question": str(q).strip(),
            "answer": strip_html(ws.cell(r, 4).value),
            "similar": sims,
        }


def build_faq_reference(entries):
    n = len(entries)
    # A token in more than ~12% of entries ("hero", "game", "match") says
    # nothing about WHICH entry a player question belongs to, and would let
    # unrelated questions look "covered" in gap detection. Drop them.
    df = collections.Counter()
    for e in entries:
        df.update(set(tokens(e["question"] + " " + " ".join(e["similar"]) + " " + e["answer"])))
    too_common = {t for t, c in df.items() if c > max(8, int(n * 0.12))}

    out = []
    for e in entries:
        q_tokens = [t for t in tokens(e["question"]) if t not in STOPWORDS and t not in too_common]
        sim_tokens = [t for s in e["similar"] for t in tokens(s) if t not in STOPWORDS and t not in too_common]
        a_tokens = [t for t in tokens(e["answer"]) if t not in STOPWORDS and t not in too_common][:8]
        combined = (e["question"] + " " + e["answer"]).lower()
        slang = []
        for formal, syns in PLAYER_SLANG.items():
            if re.search(r"\b" + re.escape(formal) + r"\b", combined):
                slang.extend(syns)
        seen, kws = set(), []
        for k in slang + q_tokens + sim_tokens + a_tokens:
            if k not in seen:
                seen.add(k)
                kws.append(k)
        out.append({"id": e["id"], "category": e["category"], "question": e["question"], "keywords": kws[:18]})
    return out


# Capitalised words in the FAQ that are ordinary English, not in-game terms.
# Matching these as "game vocabulary" would make unrelated chat look on-topic.
GENERIC_TERMS = {
    "background", "battle", "contact us", "discord", "frame", "fighter", "knockout", "lose", "menu",
    "mode", "profile", "resource", "scholar", "settings", "support", "stars", "android", "live ops",
    "immortal", "legend", "casual", "soccer knockout",
}


def extract_game_terms(entries):
    """Title-Case in-game terms (Blitz Pass, Hero Road, Demotion Shields ...)
    that appear capitalised mid-sentence in at least two entries."""
    pattern = re.compile(r"(?<![.!?]\s)(?<!^)\b([A-Z][a-zA-Z']+(?:\s+[A-Z][a-zA-Z']+){0,2})\b")
    doc_terms = collections.defaultdict(set)
    for i, e in enumerate(entries):
        text = e["question"] + " " + e["answer"]
        for m in pattern.finditer(text):
            term = m.group(1).strip()
            # a leading sentence-start word is not part of the term
            words = term.split()
            if words and words[0].lower() in STOPWORDS:
                words = words[1:]
            if not words:
                continue
            doc_terms[" ".join(words).lower()].add(i)
    terms = [t for t, docs in doc_terms.items() if len(docs) >= 2 and len(t) >= 4 and t not in GENERIC_TERMS]
    return sorted(terms)


def drop_hero_compounds(terms, hero_names_latin):
    """'cravius corrupt' style skin names add nothing beyond the hero name itself."""
    heroes = {h.lower() for h in hero_names_latin}
    return [t for t in terms if not (" " in t and t.split()[0] in heroes)]


def find_roster_row(rows, hero_hint_count=6):
    """The 'Which heroes are in the game?' entry lists every hero name in a
    fixed order, identically across all language sheets, so names can be
    paired by position to get exact localized forms."""
    for r in rows:
        paren_free = re.sub(r"\([^)]*\)", " ", r["answer"])
        words = paren_free.split()
        if 8 <= len(words) <= 16 and all(w[0].isupper() for w in words if w[0].isalpha()):
            return r["id"], words
    return None, None


def clean_name_list(answer):
    return re.sub(r"\([^)]*\)", " ", answer).split()


# Frequent first words that are articles or prepositions in their language,
# not interrogatives. They start many FAQ questions only because the next
# word is the question word ("Os heróis...", "Di mana...").
NOT_INTERROGATIVE = {"as", "os", "di", "se"}


def question_starters(rows, min_count=4):
    first = collections.Counter()
    for r in rows:
        w = re.findall(r"[^\W\d_]+", r["question"].lower())
        if w:
            first[w[0]] += 1
    return sorted(w for w, c in first.items() if c >= min_count and len(w) >= 2 and w not in NOT_INTERROGATIVE)


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    sheets = {name: list(sheet_rows(wb[name])) for name in wb.sheetnames}
    en = sheets["English"]

    faq_ref = build_faq_reference(en)

    roster_id, en_names = find_roster_row(en)
    hero_names = {}
    if roster_id:
        hero_names["English"] = en_names
        for lang, rows in sheets.items():
            if lang == "English":
                continue
            row = next((r for r in rows if r["id"] == roster_id), None)
            names = clean_name_list(row["answer"]) if row else []
            if len(names) == len(en_names):
                hero_names[lang] = names
            else:
                print(f"WARNING: roster for {lang} has {len(names)} names, expected {len(en_names)}; skipped", file=sys.stderr)

    game_terms = extract_game_terms(en)
    if hero_names.get("English"):
        game_terms = drop_hero_compounds(game_terms, hero_names["English"])

    vocab = {
        "generatedFrom": os.path.basename(path),
        "faqEntryCount": len(en),
        "gameTerms": game_terms,
        "heroNames": hero_names,
        "questionStarters": {lang: question_starters(rows) for lang, rows in sheets.items()},
    }

    os.makedirs(OUT_DIR, exist_ok=True)
    with open(os.path.join(OUT_DIR, "faqReference.json"), "w", encoding="utf-8") as f:
        json.dump(faq_ref, f, ensure_ascii=False, indent=2)
    with open(os.path.join(OUT_DIR, "prefilterVocabulary.json"), "w", encoding="utf-8") as f:
        json.dump(vocab, f, ensure_ascii=False, indent=2)

    print(f"faqReference.json: {len(faq_ref)} entries")
    print(f"prefilterVocabulary.json: {len(vocab['gameTerms'])} game terms, "
          f"hero names in {len(hero_names)} languages, starters for {len(vocab['questionStarters'])} languages")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(1)
    main(sys.argv[1])
